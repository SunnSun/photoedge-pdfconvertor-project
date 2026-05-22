import os
import pandas as pd
from pdf2docx import Converter
import pdfplumber
from pdf2image import convert_from_path
from pptx import Presentation

def convert_pdf_to_word(input_path, output_path):
    """將 PDF 轉換為 Word (.docx)"""
    try:
        cv = Converter(input_path)
        cv.convert(output_path, start=0, end=None)
        cv.close()
        return True
    except Exception as e:
        print(f"Error converting PDF to Word: {e}")
        return False

def convert_pdf_to_excel(input_path, output_path):
    try:
        with pdfplumber.open(input_path) as pdf:
            all_tables = []
            for page in pdf.pages:
                tables = page.extract_tables()
                for table in tables:
                    df = pd.DataFrame(table)
                    all_tables.append(df)
            
            if all_tables:
                with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
                    for i, df in enumerate(all_tables):
                        # each table goes to a separate sheet
                        df.to_excel(writer, sheet_name=f'Table_{i+1}', index=False, header=False)
                return True
        return False
    except Exception as e:
        print(f"Error converting PDF to Excel: {e}")
        return False

def convert_pdf_to_ppt(input_path, output_path):
    try:
        poppler_bin_path = r"C:\Program Files\poppler-26.02.0\Library\bin"
        images = convert_from_path(input_path , poppler_path=poppler_bin_path)
        prs = Presentation()
        blank_slide_layout = prs.slide_layouts[6] # 使用空白佈局
        
        for i, image in enumerate(images):

            tmp_img_path = f"{input_path}_page_{i}.png"
            image.save(tmp_img_path, 'PNG')
            
            slide = prs.slides.add_slide(blank_slide_layout)
            slide.shapes.add_picture(tmp_img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
            
            os.remove(tmp_img_path)
            
        prs.save(output_path)
        return True
    except Exception as e:
        print(f"Error converting PDF to PPT: {e}")
        return False